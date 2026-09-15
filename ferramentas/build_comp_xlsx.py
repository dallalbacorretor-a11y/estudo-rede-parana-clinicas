# -*- coding: utf-8 -*-
import os, sys, datetime
from collections import Counter
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import _local  # noqa: F401  (fixa o diretorio de trabalho)
from common import PLANOS, ORDEM_CAT, norm
from build_comp import concilia
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = _local.RAIZ
D1 = os.path.join(BASE, "01 - REDE CREDENCIADA", "Paraná Clínicas")
D2 = os.path.join(BASE, "02 - COMPARATIVOS DE REDE")
CURTOS = [p[2] for p in PLANOS]
CORES = {"400 QP": "A80A32", "600 QC": "1F6B4E", "CIM QC": "2A2ABF"}
TINT = {"400 QP": "FBEEF1", "600 QC": "EDF5F1", "CIM QC": "EEEEFB"}
NAVY = "10233F"
thin = Side(style="thin", color="D9D9D9")
BORD = Border(left=thin, right=thin, top=thin, bottom=thin)
HOJE = datetime.date.today().strftime("%d/%m/%Y")

d = concilia()
d.sort(key=lambda r: (ORDEM_CAT.index(r["categoria"]) if r["categoria"] in ORDEM_CAT else 99,
                      -len(r["planos"]), r["cidade"] != "Curitiba", r["cidade"], norm(r["nome"])))

wb = Workbook(); ws = wb.active; ws.title = "Comparativo"
COLS = ["Categoria", "Tipo", "Prestador", "Especialidades", "Cidade", "Bairro",
        "Endereço", "CEP", "Telefone"] + CURTOS
LARG = [22, 18, 42, 46, 20, 22, 44, 12, 16, 11, 11, 11]
ws.append(COLS)
for r in d:
    ws.append([r["categoria"], r["tipo"], r["nome"], r["esp"], r["cidade"], r["bairro"],
               r["endereco"], r["cep"], r["tel"]] +
              ["✔" if c in r["planos"] else "" for c in CURTOS])
for c in range(1, len(COLS) + 1):
    cell = ws.cell(row=1, column=c)
    cell.font = Font(bold=True, color="FFFFFF", size=10)
    cell.fill = PatternFill("solid", fgColor=CORES[COLS[c - 1]] if COLS[c - 1] in CORES else NAVY)
    cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
    ws.column_dimensions[get_column_letter(c)].width = LARG[c - 1]
ws.row_dimensions[1].height = 28
ws.freeze_panes = "D2"
ws.auto_filter.ref = ws.dimensions
for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.font = Font(size=9)
        cell.border = BORD
for i, k in enumerate(CURTOS):
    col = get_column_letter(10 + i)
    for r_ in range(2, ws.max_row + 1):
        c = ws["%s%d" % (col, r_)]
        c.fill = PatternFill("solid", fgColor=TINT[k])
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.font = Font(size=11, bold=True, color=CORES[k])

ws2 = wb.create_sheet("Resumo")
ws2.append(["Comparativo de rede credenciada — Paraná Clínicas"])
ws2.append(["Consulta realizada em", HOJE])
ws2.append(["Fonte", "Busca oficial de rede credenciada da Paraná Clínicas"])
ws2.append([])
ws2.append(["Plano", "Registro ANS", "Prestadores/endereços"])
for cod, nome, curto, ans in PLANOS:
    ws2.append([nome, ans, sum(1 for x in d if curto in x["planos"])])
ws2.append([])
ws2.append(["Abrangência", "Prestadores/endereços"])
cc = Counter(" + ".join(x["planos"]) for x in d)
for k, v in sorted(cc.items(), key=lambda x: -x[1]):
    ws2.append([k, v])
ws2.append(["TOTAL (rede consolidada)", len(d)])
ws2.append([])
ws2.append(["Categoria"] + CURTOS)
for cat in ORDEM_CAT:
    ws2.append([cat] + [sum(1 for x in d if x["categoria"] == cat and c in x["planos"]) for c in CURTOS])
ws2.append([])
ws2.append(["Cidade"] + CURTOS)
for cid in sorted(set(x["cidade"] for x in d),
                  key=lambda c: -sum(1 for x in d if x["cidade"] == c)):
    ws2.append([cid] + [sum(1 for x in d if x["cidade"] == cid and c in x["planos"]) for c in CURTOS])
for col, w in (("A", 46), ("B", 18), ("C", 22), ("D", 12)):
    ws2.column_dimensions[col].width = w
ws2["A1"].font = Font(bold=True, size=13, color=NAVY)
for r_ in ws2.iter_rows():
    for c in r_:
        if c.row > 1 and c.column == 1 and c.value in ("Plano", "Abrangência", "Categoria", "Cidade"):
            c.font = Font(bold=True)

fn = os.path.join(D2, "COMPARATIVO REDE PARANA CLINICAS.xlsx")
wb.save(fn); print("ok", fn, len(d))

# ---------- RESUMO.md ----------
L = []
L.append("# Rede credenciada — Paraná Clínicas\n")
L.append("**Consulta realizada em %s** na busca oficial de rede credenciada da operadora "
         "(https://www.paranaclinicas.com.br/beneficiario/rede-credenciada/).\n" % HOJE)
L.append("Operadora: Paraná Clínicas (grupo SulAmérica). Abrangência dos três planos: **Paraná — "
         "Curitiba e Região Metropolitana**.\n")
L.append("## Planos analisados\n")
L.append("| Plano | Registro ANS | Contratação | Prestadores/endereços |")
L.append("|---|---|---|---|")
for cod, nome, curto, ans in PLANOS:
    L.append("| %s | %s | Coletivo empresarial | %d |" % (nome, ans, sum(1 for x in d if curto in x["planos"])))
L.append("\n**Rede consolidada (união dos três): %d prestadores/endereços.**\n" % len(d))
L.append("## Abrangência cruzada\n")
L.append("| Atende | Prestadores/endereços |")
L.append("|---|---|")
for k, v in sorted(cc.items(), key=lambda x: -x[1]):
    L.append("| %s | %d |" % (k, v))
L.append("\n## Por categoria\n")
L.append("| Categoria | " + " | ".join(CURTOS) + " |")
L.append("|---|" + "---|" * len(CURTOS))
for cat in ORDEM_CAT:
    L.append("| %s | %s |" % (cat, " | ".join(
        str(sum(1 for x in d if x["categoria"] == cat and c in x["planos"])) for c in CURTOS)))
L.append("\n## Por cidade\n")
L.append("| Cidade | " + " | ".join(CURTOS) + " |")
L.append("|---|" + "---|" * len(CURTOS))
for cid in sorted(set(x["cidade"] for x in d), key=lambda c: -sum(1 for x in d if x["cidade"] == c)):
    L.append("| %s | %s |" % (cid, " | ".join(
        str(sum(1 for x in d if x["cidade"] == cid and c in x["planos"])) for c in CURTOS)))
L.append("\n## Observações\n")
L.append("- As **unidades próprias CIM** (Centros Integrados de Medicina) aparecem como tipo de prestador "
         "separado: CIM Água Verde, CIM CIC, CIM São José dos Pinhais e CIM Araucária. Os profissionais que "
         "atendem nessas unidades são listados individualmente pela operadora.")
L.append("- O plano **Paraná CIM AHO QC** tem rede mais restrita: aparece em **10 cidades**, contra 17 dos "
         "planos 400 e 600.")
L.append("- O plano **Paraná 600 AHO QC** é o de rede mais ampla dos três.")
L.append("- Os planos **400 QC** (ANS 507886265) e **600 QP** (ANS 507889260) são as variantes de acomodação "
         "dos mesmos produtos e compartilham a mesma rede — por isso não foram tabulados separadamente.")
L.append("\n> A rede credenciada é definida e alterada exclusivamente pela operadora. "
         "**Confirme no portal da Paraná Clínicas antes de contratar.**\n")
L.append("_Mazza Broker · Alan Vinicius Dall Alba · (41) 99547-6715 · alan.vinicius@mazzabroker.com.br_")
fn = os.path.join(D1, "RESUMO.md")
open(fn, "w", encoding="utf-8").write("\n".join(L))
print("ok", fn)
