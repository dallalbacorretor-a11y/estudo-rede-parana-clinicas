# -*- coding: utf-8 -*-
import os, sys, json
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import _local  # noqa: F401  (fixa o diretorio de trabalho)
from common import PLANOS, ORDEM_CAT, carrega, norm
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = _local.RAIZ
D1 = os.path.join(BASE, "01 - REDE CREDENCIADA", "Paraná Clínicas")
D2 = os.path.join(BASE, "02 - COMPARATIVOS DE REDE")
VERDE = "1F6B4E"; CINZA = "F2F2F2"
thin = Side(style="thin", color="D9D9D9")
BORD = Border(left=thin, right=thin, top=thin, bottom=thin)

def estiliza(ws, ncols, larguras):
    for c in range(1, ncols+1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor=VERDE)
        cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = larguras[c-1]
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.font = Font(size=9)
            cell.border = BORD

COLS = ["Categoria", "Tipo", "Prestador", "Especialidades", "Cidade", "Bairro",
        "Endereço", "CEP", "Telefone", "Telefone 2"]
LARG = [22, 18, 42, 46, 20, 22, 44, 12, 16, 16]

def sheet_rede(ws, dados):
    ws.append(COLS)
    dados = sorted(dados, key=lambda r: (ORDEM_CAT.index(r["categoria"]) if r["categoria"] in ORDEM_CAT else 99,
                                         r["cidade"], norm(r["nome"]), r["bairro"]))
    for r in dados:
        ws.append([r["categoria"], r["tipo"], r["nome"], r["esp"], r["cidade"], r["bairro"],
                   r["endereco"], r["cep"], r["tel"], r["tel2"]])
    estiliza(ws, len(COLS), LARG)

# ---- 1) um xlsx por plano
todos = {}
for cod, nome, curto, ans in PLANOS:
    dados = carrega(cod)
    todos[cod] = dados
    wb = Workbook(); ws = wb.active; ws.title = "Rede"
    sheet_rede(ws, dados)
    ws2 = wb.create_sheet("Resumo")
    ws2.append(["Plano", nome]); ws2.append(["Registro ANS", ans])
    ws2.append(["Código do plano (busca)", cod]); ws2.append([])
    ws2.append(["Categoria", "Prestadores/endereços"])
    from collections import Counter
    cc = Counter(r["categoria"] for r in dados)
    for c in ORDEM_CAT:
        if cc.get(c): ws2.append([c, cc[c]])
    ws2.append(["TOTAL", len(dados)])
    ws2.append([])
    ws2.append(["Cidade", "Prestadores/endereços"])
    cd = Counter(r["cidade"] for r in dados)
    for c, n in sorted(cd.items(), key=lambda x: -x[1]): ws2.append([c, n])
    for col, w in (("A", 34), ("B", 24)): ws2.column_dimensions[col].width = w
    for cel in ("A1", "A2", "A3", "A5", "B5"): ws2[cel].font = Font(bold=True)
    fn = os.path.join(D1, "REDE %s.xlsx" % nome.upper().replace("PARANÁ ", "PARANA "))
    wb.save(fn); print("ok", fn, len(dados))

json.dump({k: v for k, v in todos.items()}, open("dados/todos.json", "w", encoding="utf-8"), ensure_ascii=False)
